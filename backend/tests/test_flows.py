"""End-to-end API flows A–H, account lifecycle, audit, search privacy and document validation.

Every test drives the real HTTP API (TestClient) against a freshly seeded synthetic database, performs actions
as the CURRENT logged-in user and checks the effects on access, eligibility, alerts and the audit log.
"""
from __future__ import annotations

import io
from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.core.config import settings
from app.db import SessionLocal
from app.models import Person
from app.seed import DEMO_PASSWORD, seed

JAIL_A = "Central Prison Parappana Agrahara"
JAIL_B = "Central Prison Mysuru"


@pytest.fixture(scope="module")
def client():
    settings.demo_mode = True
    seed(population=4, reset=True, verbose=False)  # full demo set: the two headline prisoners + edge cases
    from app.main import app
    with TestClient(app) as c:
        yield c
    settings.demo_mode = False


def login(client, email, password=DEMO_PASSWORD):
    r = client.post("/api/auth/login", json={"email": email, "password": password})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def pid(truth: str) -> str:
    with SessionLocal() as db:
        return db.scalar(select(Person.id).where(Person.synthetic_truth_id == truth))


def audit_actions(client, admin, **params) -> list[dict]:
    return client.get("/api/admin/audit", params={"limit": 1000, **params}, headers=admin).json()


def verify_intake(client, person_id: str) -> None:
    """The reviewer checks a new prisoner's entered details (the step before a lawyer can be assigned)."""
    reviewer = login(client, "reviewer@nyayasetu.test")
    item = next(i for i in client.get("/api/review", params={"kind": "new_prisoner"}, headers=reviewer).json()
                if i["ref_id"] == person_id)
    assert client.post(f"/api/review/{item['id']}/resolve", json={"action": "confirm"}, headers=reviewer).status_code == 200


def new_prisoner(client, headers, name="Kiran Test", jail=None, district=None, **case):
    today = settings.today()
    body = {"name": name, "relative_name": "Ramappa", "relation": "s/o", "gender": "male", "jail": jail, "district": district,
            "case": {"court": "JMFC Test Court", "case_number": "CC 900/2026", "fir_number": "12/2026",
                     "charges": [{"act": "BNS", "section": "303(2)"}],
                     "offence_date": (today - timedelta(days=45)).isoformat(),
                     "arrest_date": (today - timedelta(days=40)).isoformat(),
                     "first_remand_date": (today - timedelta(days=39)).isoformat(), **case}}
    return client.post("/api/persons", json=body, headers=headers)


# ------------------------------------------------------------------ the two headline demo prisoners
def test_headline_demo_prisoners(client):
    lawyer = login(client, "lawyer@nyayasetu.test")
    ravi = client.get(f"/api/persons/{pid('T-DEMO-NOT')}", headers=lawyer).json()
    suresh = client.get(f"/api/persons/{pid('T-DEMO-ELIG')}", headers=lawyer).json()
    rc, sc = ravi["eligibility"]["cases"][0], suresh["eligibility"]["cases"][0]
    assert ravi["name"] == "Ravi Kumar" and ravi["status"] == "NOT_YET" and ravi["demo_label"].startswith("DEMO A")
    assert (rc["custody_days"], rc["accused_days"], rc["counted_days"], rc["threshold_days"]) == (263, 21, 242, 366)
    assert rc["projected_date"] == (settings.today() + timedelta(days=124)).isoformat()
    assert suresh["name"] == "Suresh Kumar" and suresh["status"] == "ELIGIBLE"
    assert (sc["custody_days"], sc["accused_days"], sc["counted_days"], sc["threshold_days"]) == (991, 30, 961, 853)
    assert sc["days_overdue"] == 108 and sc["eligible_from_date"] == (settings.today() - timedelta(days=108)).isoformat()
    # every document of both prisoners agrees with the record: no document/record conflicts
    for p in (ravi, suresh):
        assert not any(f["code"] == "DOCUMENT_RECORD_CONFLICT" for f in p["eligibility"]["cases"][0]["flags"])
        assert not p["eligibility"]["cases"][0]["needs_verification"]
    # the headline name is unique in the demo data (the old eligible "Ravi Kumar" is now Ramesh Babu)
    admin = login(client, "admin@nyayasetu.test")
    assert [r["name"] for r in client.get("/api/register", headers=admin).json()].count("Ravi Kumar") == 1


# ------------------------------------------------------------------ FLOW A: new prisoner → unassigned → assigned
def test_flow_a_new_prisoner_starts_unassigned(client):
    jail, lawyer, dlsa = login(client, "jail@nyayasetu.test"), login(client, "lawyer@nyayasetu.test"), login(client, "dlsa@nyayasetu.test")
    admin = login(client, "admin@nyayasetu.test")
    r = new_prisoner(client, jail, name="Flow A Prisoner")
    assert r.status_code == 200, r.text
    new_id = r.json()["id"]
    assert r.json()["assigned_lawyer_id"] is None and r.json()["status"] in ("NOT_YET", "REVIEW")
    assert new_id not in {x["id"] for x in client.get("/api/persons", headers=lawyer).json()}
    assert client.get(f"/api/persons/{new_id}", headers=lawyer).status_code == 404
    row = next(x for x in client.get("/api/persons", headers=dlsa).json() if x["id"] == new_id)  # same district
    assert row["assigned_lawyer_id"] is None
    lawyer_id = client.get("/api/auth/me", headers=lawyer).json()["id"]
    # not assignable until the reviewer has verified the entered details
    assert client.post(f"/api/persons/{new_id}/assign", json={"lawyer_id": lawyer_id}, headers=admin).status_code == 409
    assert any(r["id"] == new_id and r["intake"] == "pending_review" for r in client.get("/api/register", headers=admin).json())
    verify_intake(client, new_id)
    assert client.post(f"/api/persons/{new_id}/assign", json={"lawyer_id": lawyer_id}, headers=admin).status_code == 200
    assert client.get(f"/api/persons/{new_id}", headers=lawyer).status_code == 200
    acts = {(a["action"], a["entity_id"]) for a in audit_actions(client, admin, entity_id=new_id)}
    assert {("create_prisoner", new_id), ("verify_intake", new_id), ("assign_lawyer", new_id)} <= acts


def test_new_prisoner_validation_and_jail_scope(client):
    jail = login(client, "jail@nyayasetu.test")
    today = settings.today()
    assert new_prisoner(client, jail, jail=JAIL_B).status_code == 403  # jail staff only add to their own jail
    bad = new_prisoner(client, jail, arrest_date=(today + timedelta(days=3)).isoformat())
    assert bad.status_code == 422 and "future" in bad.json()["detail"]
    bad = new_prisoner(client, jail, first_remand_date=(today - timedelta(days=60)).isoformat())
    assert bad.status_code == 422 and "before the arrest" in bad.json()["detail"]
    assert client.post("/api/persons", json={"name": "X"}, headers=jail).status_code == 422
    assert new_prisoner(client, login(client, "lawyer@nyayasetu.test")).status_code == 403


def test_new_prisoner_with_unknown_section_goes_to_review(client):
    jail = login(client, "jail@nyayasetu.test")
    r = new_prisoner(client, jail, name="Unknown Section", charges=[{"act": "IPC", "section": "9999"}])
    assert r.status_code == 200 and r.json()["status"] == "REVIEW"


# ------------------------------------------------------------------ FLOW B: new lawyer
def test_flow_b_new_lawyer_starts_with_zero(client):
    dlsa, admin = login(client, "dlsa@nyayasetu.test"), login(client, "admin@nyayasetu.test")
    r = client.post("/api/lawyers", json={"email": "new.lawyer@nyayasetu.test", "name": "Adv. New Lawyer",
                                          "password": "temporary-pass-1"}, headers=dlsa)
    assert r.status_code == 200 and r.json()["status"] == "pending_approval"
    uid = r.json()["id"]
    assert client.post("/api/auth/login", json={"email": "new.lawyer@nyayasetu.test", "password": "temporary-pass-1"}).status_code == 401
    assert client.post(f"/api/lawyers/{uid}/approve", headers=dlsa).status_code == 200
    new = login(client, "new.lawyer@nyayasetu.test", "temporary-pass-1")
    assert client.get("/api/persons", headers=new).json() == []
    target = pid("T-DEMO-NOT")
    assert client.post(f"/api/persons/{target}/assign", json={"lawyer_id": uid}, headers=dlsa).status_code == 200
    assert [x["id"] for x in client.get("/api/persons", headers=new).json()] == [target]
    acts = [a["action"] for a in audit_actions(client, admin, entity_id=uid)]
    assert "create_user" in acts and "approve_user" in acts
    # give Ravi back to the demo lawyer for the other tests
    lawyer_id = client.get("/api/auth/me", headers=login(client, "lawyer@nyayasetu.test")).json()["id"]
    client.post(f"/api/persons/{target}/assign", json={"lawyer_id": lawyer_id}, headers=dlsa)
    assert client.get(f"/api/persons/{target}", headers=new).status_code == 404  # immediately loses access


# ------------------------------------------------------------------ FLOW C: lawyer leaves
def test_flow_c_lawyer_leaves(client):
    dlsa, admin = login(client, "dlsa@nyayasetu.test"), login(client, "admin@nyayasetu.test")
    r = client.post("/api/lawyers", json={"email": "leaving@nyayasetu.test", "name": "Adv. Leaving",
                                          "password": "temporary-pass-2", "approve": True}, headers=dlsa)
    uid = r.json()["id"]
    leaving = login(client, "leaving@nyayasetu.test", "temporary-pass-2")
    created = new_prisoner(client, login(client, "jail@nyayasetu.test"), name="Flow C Prisoner").json()["id"]
    verify_intake(client, created)
    client.post(f"/api/persons/{created}/assign", json={"lawyer_id": uid}, headers=dlsa)
    assert client.get(f"/api/persons/{created}", headers=leaving).status_code == 200
    r = client.post(f"/api/lawyers/{uid}/deactivate", headers=dlsa)
    assert r.status_code == 200 and r.json()["unassigned_prisoners"] == 1
    assert client.get(f"/api/persons/{created}", headers=leaving).status_code == 401  # old token dead at once
    assert client.post("/api/auth/login", json={"email": "leaving@nyayasetu.test", "password": "temporary-pass-2"}).status_code == 401
    row = next(x for x in client.get("/api/persons", headers=dlsa).json() if x["id"] == created)
    assert row["assigned_lawyer_id"] is None
    assert any(a["kind"] == "LAWYER_REASSIGNMENT_NEEDED" and a["person_id"] == created
               for a in client.get("/api/alerts", headers=dlsa).json())
    lawyer2 = login(client, "lawyer2@nyayasetu.test")
    l2 = client.get("/api/auth/me", headers=lawyer2).json()["id"]
    assert client.post(f"/api/persons/{created}/assign", json={"lawyer_id": l2}, headers=dlsa).status_code == 200
    assert client.get(f"/api/persons/{created}", headers=lawyer2).status_code == 200
    assert client.post(f"/api/persons/{created}/assign", json={"lawyer_id": uid}, headers=dlsa).status_code == 400  # inactive
    client.post(f"/api/lawyers/{uid}/activate", headers=dlsa)  # reactivation does NOT restore assignments
    back = login(client, "leaving@nyayasetu.test", "temporary-pass-2")
    assert client.get("/api/persons", headers=back).json() == []
    acts = [a["action"] for a in audit_actions(client, admin, entity_id=created)]
    assert "unassign_lawyer" in acts and "assign_lawyer" in acts


# ------------------------------------------------------------------ reassignment and drafts
def test_reassignment_moves_access_and_keeps_data(client):
    lawyer, lawyer2, dlsa = login(client, "lawyer@nyayasetu.test"), login(client, "lawyer2@nyayasetu.test"), login(client, "dlsa@nyayasetu.test")
    jail = login(client, "jail@nyayasetu.test")
    target = pid("T-DEMO-ELIG")
    before = client.get(f"/api/persons/{target}", headers=lawyer).json()
    case_id = before["cases"][0]["id"]
    draft = client.post(f"/api/persons/{target}/drafts", json={"case_id": case_id, "type": "section_479", "use_llm": False},
                        headers=lawyer).json()
    # jail staff (same jail) must not see the lawyer's draft, nor export it
    assert draft["id"] not in {d["id"] for d in client.get(f"/api/persons/{target}/drafts", headers=jail).json()}
    assert client.get(f"/api/drafts/{draft['id']}", headers=jail).status_code == 404
    assert client.get(f"/api/drafts/{draft['id']}/export?format=pdf", headers=jail).status_code == 404
    l1 = client.get("/api/auth/me", headers=lawyer).json()["id"]
    l2 = client.get("/api/auth/me", headers=lawyer2).json()["id"]
    assert client.post(f"/api/persons/{target}/assign", json={"lawyer_id": l2}, headers=dlsa).status_code == 200
    for path in (f"/api/persons/{target}", f"/api/persons/{target}/insights", f"/api/drafts/{draft['id']}",
                 f"/api/drafts/{draft['id']}/export?format=docx"):
        assert client.get(path, headers=lawyer).status_code == 404, path
        assert client.get(path, headers=lawyer2).status_code == 200, path
    after = client.get(f"/api/persons/{target}", headers=lawyer2).json()
    assert after["status"] == before["status"] and len(after["documents"]) == len(before["documents"])
    assert after["eligibility"]["cases"][0]["counted_days"] == before["eligibility"]["cases"][0]["counted_days"]
    client.post(f"/api/persons/{target}/assign", json={"lawyer_id": l1}, headers=dlsa)


# ------------------------------------------------------------------ FLOW D: a new document changes the result
def test_flow_d_release_document_sends_to_review_until_confirmed(client):
    lawyer, reviewer, admin = login(client, "lawyer@nyayasetu.test"), login(client, "reviewer@nyayasetu.test"), login(client, "admin@nyayasetu.test")
    target = pid("T-ELIG")
    p = client.get(f"/api/persons/{target}", headers=lawyer).json()
    assert p["status"] == "ELIGIBLE"
    rel = (settings.today() - timedelta(days=20)).strftime("%d/%m/%Y")
    text = f"RELEASE ORDER\nIn the Court of JMFC\nAccused: {p['name']}\nThe accused is set at liberty.\nDate of release: {rel}\n"
    r = client.post(f"/api/persons/{target}/documents", files={"file": ("release.txt", text.encode())}, headers=lawyer)
    assert r.status_code == 200 and r.json()["doc_type"] == "release_order"
    p2 = client.get(f"/api/persons/{target}", headers=lawyer).json()
    c = p2["eligibility"]["cases"][0]
    assert p2["status"] == "REVIEW" and any(f["code"] == "DOCUMENT_RECORD_CONFLICT" for f in c["flags"])
    assert any(a["kind"] == "RESULT_CHANGED" and a["person_id"] == target for a in client.get("/api/alerts", headers=lawyer).json())
    # the reviewer rejects the fact (e.g. wrong prisoner's order) → it is no longer used, the result returns
    fact = next(f for f in p2["facts"] if f["field"] == "release_date")
    assert client.patch(f"/api/facts/{fact['id']}", json={"action": "reject", "note": "not this prisoner"}, headers=reviewer).status_code == 200
    assert client.get(f"/api/persons/{target}", headers=lawyer).json()["status"] == "ELIGIBLE"
    acts = {a["action"] for a in audit_actions(client, admin)}
    assert {"upload_document", "fact_reject"} <= acts


# ------------------------------------------------------------------ FLOW E: transfer
def test_flow_e_transfer_moves_jail_scope_not_eligibility(client):
    jail, lawyer, admin = login(client, "jail@nyayasetu.test"), login(client, "lawyer@nyayasetu.test"), login(client, "admin@nyayasetu.test")
    target = pid("T-DEMO-ELIG")
    before = client.get(f"/api/persons/{target}", headers=lawyer).json()
    r = client.post(f"/api/persons/{target}/transfer", json={"to_jail": JAIL_B, "date": (settings.today() - timedelta(days=5)).isoformat()},
                    headers=jail)
    assert r.status_code == 200, r.text
    after = client.get(f"/api/persons/{target}", headers=lawyer).json()  # lawyer access unchanged
    assert after["jail"] == JAIL_B and after["district"] == "Mysuru"
    assert after["status"] == before["status"]
    for k in ("custody_days", "counted_days", "eligible_from_date", "days_overdue"):
        assert after["eligibility"]["cases"][0][k] == before["eligibility"]["cases"][0][k], k  # no gap, no double count
    assert any(e["type"] == "transfer" and e["jail"] == JAIL_B for e in after["custody_events"])
    assert client.get(f"/api/persons/{target}", headers=jail).status_code == 404  # old jail loses access
    client.post("/api/admin/users", json={"email": "jail.mysuru@nyayasetu.test", "name": "Mysuru jail", "role": "jail_staff",
                                          "password": "mysuru-pass-123", "jail": JAIL_B}, headers=admin)
    mysuru = login(client, "jail.mysuru@nyayasetu.test", "mysuru-pass-123")
    assert client.get(f"/api/persons/{target}", headers=mysuru).status_code == 200  # new jail gains access
    assert "transfer_prisoner" in [a["action"] for a in audit_actions(client, admin, entity_id=target)]
    client.post(f"/api/persons/{target}/transfer", json={"to_jail": JAIL_A, "to_district": "Bengaluru Urban",
                                                         "date": settings.today().isoformat()}, headers=mysuru)


# ------------------------------------------------------------------ FLOW F: release
def test_flow_f_release_closes_custody_and_resolves_alerts(client):
    jail, lawyer, dlsa, admin = (login(client, e) for e in ("jail@nyayasetu.test", "lawyer@nyayasetu.test",
                                                            "dlsa@nyayasetu.test", "admin@nyayasetu.test"))
    target = pid("T-CRIT")
    before = client.get(f"/api/persons/{target}", headers=lawyer).json()
    assert before["urgency"] == "CRITICAL_MUST_RELEASE"
    open_before = [a for a in client.get("/api/alerts", headers=lawyer).json() if a["person_id"] == target]
    assert any(a["kind"] == "CRITICAL_MUST_RELEASE" for a in open_before)
    day = settings.today() - timedelta(days=2)
    assert client.post(f"/api/persons/{target}/release", json={"date": (day - timedelta(days=5000)).isoformat()}, headers=jail).status_code == 422
    r = client.post(f"/api/persons/{target}/release", json={"date": day.isoformat(), "note": "released on court order"}, headers=jail)
    assert r.status_code == 200 and r.json()["status"] == "NOT_APPLICABLE"
    assert client.post(f"/api/persons/{target}/release", json={"date": day.isoformat()}, headers=jail).status_code == 409
    for h in (lawyer, dlsa, jail):  # everyone in scope sees the same updated status
        p = client.get(f"/api/persons/{target}", headers=h).json()
        assert p["status"] == "NOT_APPLICABLE" and p["eligibility"]["cases"][0]["custody_days"] == before["eligibility"]["cases"][0]["custody_days"] - 2
    assert not any(a["kind"] == "CRITICAL_MUST_RELEASE" and a["person_id"] == target for a in client.get("/api/alerts", headers=lawyer).json())
    r1 = client.post(f"/api/persons/{target}/recompute", headers=lawyer).json()
    r2 = client.post(f"/api/persons/{target}/recompute", headers=lawyer).json()
    assert r1["status"] == r2["status"] == "NOT_APPLICABLE" and r2["changed"] is False
    assert next(r for r in client.get("/api/register", headers=admin).json() if r["id"] == target)["status"] == "NOT_APPLICABLE"
    assert "record_release" in [a["action"] for a in audit_actions(client, admin, entity_id=target)]


# ------------------------------------------------------------------ FLOW G: unauthorized lawyer
def test_flow_g_unauthorized_lawyer_gets_nothing(client):
    lawyer, lawyer2, admin = login(client, "lawyer@nyayasetu.test"), login(client, "lawyer2@nyayasetu.test"), login(client, "admin@nyayasetu.test")
    target = pid("T-DEMO-NOT")
    p = client.get(f"/api/persons/{target}", headers=lawyer).json()
    doc_id, fact_id, case_id = p["documents"][0]["id"], p["facts"][0]["id"], p["cases"][0]["id"]
    missing = client.get("/api/persons/zzzzzzzzzzzz", headers=lawyer2)
    attempts = [
        client.get(f"/api/persons/{target}", headers=lawyer2),
        client.get(f"/api/persons/{target}/insights", headers=lawyer2),
        client.get(f"/api/persons/{target}/drafts", headers=lawyer2),
        client.get(f"/api/documents/{doc_id}", headers=lawyer2),
        client.post(f"/api/persons/{target}/recompute", headers=lawyer2),
        client.post(f"/api/persons/{target}/drafts", json={"case_id": case_id, "type": "section_479", "use_llm": False}, headers=lawyer2),
        client.patch(f"/api/facts/{fact_id}", json={"action": "confirm"}, headers=lawyer2),
        client.post(f"/api/persons/{target}/documents", files={"file": ("x.txt", b"FIR")}, headers=lawyer2),
    ]
    for r in attempts:
        assert r.status_code == 404, (r.request.url, r.status_code)
        assert r.json() == missing.json()  # identical to a made-up id: existence is not revealed
        assert "Ravi" not in r.text
    assert target not in {x["id"] for x in client.get("/api/persons", params={"q": "Ravi"}, headers=lawyer2).json()}
    assert client.get("/api/persons", params={"q": "Ravi"}, headers=lawyer2).json() == []
    assert client.get("/api/admin/audit", headers=lawyer2).status_code == 403
    denied = [a for a in audit_actions(client, admin, entity_id=target) if a["outcome"] == "denied"]
    assert len(denied) >= 5 and all(a["role"] == "legal_aid_lawyer" for a in denied)


def test_pagination_and_filters_never_widen_scope(client):
    lawyer2 = login(client, "lawyer2@nyayasetu.test")
    everything = {x["id"] for x in client.get("/api/persons", headers=lawyer2).json()}
    paged = set()
    for off in range(0, 50, 2):
        paged |= {x["id"] for x in client.get("/api/persons", params={"limit": 2, "offset": off}, headers=lawyer2).json()}
    assert paged == everything
    for st in ("ELIGIBLE", "NOT_YET", "CRITICAL_MUST_RELEASE", "REVIEW"):
        assert {x["id"] for x in client.get("/api/persons", params={"status": st}, headers=lawyer2).json()} <= everything
    assert client.get("/api/persons", params={"limit": 0}, headers=lawyer2).status_code == 422


# ------------------------------------------------------------------ FLOW H: fake / misleading documents
def _fake_pdf_bytes() -> bytes:
    from pypdf import PdfWriter
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def test_flow_h_filename_never_decides_the_document_type(client):
    lawyer = login(client, "lawyer@nyayasetu.test")
    target = pid("T-DEMO-ELIG")
    before = client.get(f"/api/persons/{target}", headers=lawyer).json()
    nonsense = b"Lorem ipsum dolor sit amet, grocery list: rice, dal, onions. Cricket score 245/6.\n"
    a = client.post(f"/api/persons/{target}/documents", files={"file": ("FIR.pdf", nonsense)}, headers=lawyer).json()
    b = client.post(f"/api/persons/{target}/documents", files={"file": ("random.txt", nonsense + b" ")}, headers=lawyer).json()
    assert a["doc_type"] == b["doc_type"] == "unknown"  # same content → same classification, whatever the name
    assert "EXTENSION_MISMATCH" in a["warnings"]
    fake = b"This is an FIR.\nTrust me.\n"
    c = client.post(f"/api/persons/{target}/documents", files={"file": ("fir_real.txt", fake)}, headers=lawyer).json()
    doc = client.get(f"/api/documents/{c['id']}", headers=lawyer).json()
    assert c["doc_type"] in ("fir", "unknown") and doc["doc_type_confidence"] < 0.8
    after = client.get(f"/api/persons/{target}", headers=lawyer).json()
    assert after["status"] == before["status"]  # no false legal calculation from junk
    kinds = {i["kind"] for i in client.get("/api/review", headers=login(client, "reviewer@nyayasetu.test")).json()}
    assert "document_type" in kinds


def test_document_for_another_prisoner_or_case_is_held_for_review(client):
    lawyer = login(client, "lawyer@nyayasetu.test")
    target = pid("T-DEMO-ELIG")
    p = client.get(f"/api/persons/{target}", headers=lawyer).json()
    other = (settings.today() - timedelta(days=3000)).strftime("%d/%m/%Y")
    wrong_person = (f"ARREST MEMO\nFIR No.: 55/2019   Police Station: X PS\nName of arrested person: Mohammed Rafiq s/o Yusuf Khan\n"
                    f"Date and time of arrest: {other} 10:00 hrs\nGrounds of arrest communicated: Yes\n")
    r = client.post(f"/api/persons/{target}/documents", files={"file": ("arrest.txt", wrong_person.encode())}, headers=lawyer).json()
    assert "NAME_MISMATCH" in r["warnings"]
    d = client.get(f"/api/documents/{r['id']}", headers=lawyer).json()
    assert all(f["confidence"] <= 0.5 for f in d["facts"])  # held, not trusted
    wrong_case = (f"ORDER ON REMAND\nIn the Court of JMFC\nCrime No. 999/2031 of Y PS\nAccused: {p['name']}\n"
                  "FIR No.: 999/2031\nRemanded to judicial custody from 01/01/2025\n")
    r = client.post(f"/api/persons/{target}/documents", files={"file": ("remand.txt", wrong_case.encode())},
                    data={"case_id": p["cases"][0]["id"]}, headers=lawyer).json()
    assert "CASE_MISMATCH" in r["warnings"]
    r = client.post(f"/api/persons/{target}/documents", files={"file": ("x.txt", b"hello")},
                    data={"case_id": "not-this-prisoners-case"}, headers=lawyer)
    assert r.status_code == 422


def test_upload_validation_fails_safely(client, monkeypatch):
    lawyer = login(client, "lawyer@nyayasetu.test")
    target = pid("T-DEMO-NOT")
    up = lambda name, data: client.post(f"/api/persons/{target}/documents", files={"file": (name, data)}, headers=lawyer)  # noqa: E731
    assert up("empty.pdf", b"").status_code == 422
    assert up("bad.exe", b"MZ....").status_code == 422
    r = up("broken.pdf", b"%PDF-1.4\n garbage without structure")
    assert r.status_code == 422 and "damaged" in r.json()["detail"]
    r = up("broken.docx", b"PK\x03\x04 word/ not really a zip")
    assert r.status_code == 422
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter()
    w.append(PdfReader(io.BytesIO(_fake_pdf_bytes())))
    w.encrypt("secret")
    buf = io.BytesIO()
    w.write(buf)
    r = up("locked.pdf", buf.getvalue())
    assert r.status_code == 422 and "password" in r.json()["detail"]
    r = up("blank.pdf", _fake_pdf_bytes())
    assert r.status_code == 200  # stored, flagged blank, sent to review — not a crash
    monkeypatch.setattr(settings, "max_upload_mb", 0)
    assert up("big.txt", b"x" * 10).status_code == 422
    monkeypatch.setattr(settings, "max_upload_mb", 25)
    import docx
    d = docx.Document()
    d.add_paragraph("JAIL ADMISSION REGISTER EXTRACT")
    d.add_paragraph("UTP No.: 7701")
    buf = io.BytesIO()
    d.save(buf)
    assert up("jail.docx", buf.getvalue()).status_code == 200
    assert up("jail.docx", buf.getvalue()).json()["duplicate"] is True


# ------------------------------------------------------------------ multiple cases
def test_closing_one_case_leaves_the_other_open(client):
    jail, admin, dlsa = login(client, "jail@nyayasetu.test"), login(client, "admin@nyayasetu.test"), login(client, "dlsa@nyayasetu.test")
    target = pid("T-MULTI")
    p = client.get(f"/api/persons/{target}", headers=dlsa).json()
    c1, c2 = p["case_details"][0], p["case_details"][1]
    r = client.patch(f"/api/cases/{c1['id']}", json={"status": "acquitted"}, headers=jail)
    assert r.status_code == 422  # an acquittal needs its date
    r = client.patch(f"/api/cases/{c1['id']}", json={"status": "acquitted", "acquittal_date": settings.today().isoformat()}, headers=jail)
    assert r.status_code == 200
    p2 = client.get(f"/api/persons/{target}", headers=dlsa).json()
    by_id = {c["id"]: c for c in p2["case_details"]}
    assert by_id[c1["id"]]["status"] == "acquitted" and by_id[c2["id"]]["status"] == c2["status"]
    assert "update_case" in [a["action"] for a in audit_actions(client, admin, entity_id=c1["id"])]


# ------------------------------------------------------------------ accounts, sessions, audit
def test_logout_revokes_the_token(client):
    h = login(client, "reviewer@nyayasetu.test")
    assert client.get("/api/review", headers=h).status_code == 200
    assert client.post("/api/auth/logout", headers=h).status_code == 200
    assert client.get("/api/review", headers=h).status_code == 401


def test_repeated_failed_logins_lock_the_account_and_are_audited(client):
    admin = login(client, "admin@nyayasetu.test")
    email = "lock.me@nyayasetu.test"
    client.post("/api/admin/users", json={"email": email, "name": "Lock Me", "role": "reviewer", "password": "correct-pass-1"}, headers=admin)
    for _ in range(5):
        assert client.post("/api/auth/login", json={"email": email, "password": "wrong"}).status_code == 401
    assert client.post("/api/auth/login", json={"email": email, "password": "correct-pass-1"}).status_code == 429
    acts = [a["action"] for a in audit_actions(client, admin)]
    assert acts.count("login_failed") >= 5 and "login_locked" in acts
    uid = next(u["id"] for u in client.get("/api/admin/users", headers=admin).json() if u["email"] == email)
    assert client.post(f"/api/admin/users/{uid}/password", json={"password": "a-brand-new-pass"}, headers=admin).status_code == 200


def test_admin_password_reset_changes_the_password_and_ends_old_sessions(client):
    import time
    admin = login(client, "admin@nyayasetu.test")
    old_session = login(client, "lawyer2@nyayasetu.test")
    assert client.get("/api/persons", headers=old_session).status_code == 200
    uid = client.get("/api/auth/me", headers=old_session).json()["id"]
    time.sleep(1.1)  # token issue times are whole seconds
    assert client.post(f"/api/admin/users/{uid}/password", json={"password": "short"}, headers=admin).status_code == 422
    assert client.post(f"/api/admin/users/{uid}/password", json={"password": "farhan-new-pass-1"}, headers=admin).status_code == 200
    assert client.get("/api/persons", headers=old_session).status_code == 401           # open session signed out
    assert client.post("/api/auth/login", json={"email": "lawyer2@nyayasetu.test", "password": DEMO_PASSWORD}).status_code == 401
    new = login(client, "lawyer2@nyayasetu.test", "farhan-new-pass-1")                   # the new password works
    assert client.get("/api/persons", headers=new).status_code == 200
    assert "reset_password" in [a["action"] for a in audit_actions(client, admin, entity_id=uid)]
    client.post(f"/api/admin/users/{uid}/password", json={"password": DEMO_PASSWORD}, headers=admin)  # restore for other tests


def test_current_session_audit_entries_are_visible_to_admin(client):
    admin_login = client.post("/api/auth/login", json={"email": "admin@nyayasetu.test", "password": DEMO_PASSWORD}).json()
    admin = {"Authorization": f"Bearer {admin_login['access_token']}"}
    since = admin_login["user"]["session_audit_id"]
    target = pid("T-DEMO-ELIG")
    client.get("/api/register", headers=admin)
    client.post("/api/admin/run-nightly", headers=admin)
    client.get(f"/api/persons/{target}", headers=admin)  # refused: case details are not the admin's
    rows = client.get("/api/admin/audit", params={"after_id": since - 1}, headers=admin).json()
    acts = {r["action"] for r in rows}
    assert {"login", "list", "run_nightly", "access_denied"} <= acts
    first = rows[-1]
    assert first["action"] == "login" and first["user_name"] and first["role"] == "system_admin"
    assert first["ts"].endswith("+00:00")  # timezone-aware: shown correctly in the browser's local time
    assert all("password" not in (r["detail"] or "").lower() for r in rows)


def test_deactivated_admin_token_rejected_and_self_deactivation_blocked(client):
    admin = login(client, "admin@nyayasetu.test")
    me = client.get("/api/auth/me", headers=admin).json()["id"]
    assert client.patch(f"/api/admin/users/{me}", json={"active": False}, headers=admin).status_code == 409


def test_role_matrix_on_admin_endpoints(client):
    roles = {e: login(client, f"{e}@nyayasetu.test") for e in ("lawyer", "jail", "dlsa", "reviewer")}
    for name, h in roles.items():
        assert client.get("/api/register", headers=h).status_code == 403, name
        assert client.get("/api/admin/audit", headers=h).status_code == 403, name
        assert client.get("/api/admin/users", headers=h).status_code == 403, name
        assert client.post("/api/admin/run-nightly", headers=h).status_code == 403, name
    assert client.post("/api/lawyers", json={"email": "x@y.z", "name": "X", "password": "0123456789"}, headers=roles["jail"]).status_code == 403
    assert client.get("/api/persons", headers=roles["reviewer"]).status_code == 403
    assert client.post("/api/resolution/scan", headers=roles["lawyer"]).status_code == 403
    assert client.post(f"/api/persons/{pid('T-DEMO-NOT')}/transfer", json={"to_jail": JAIL_B, "date": date.today().isoformat()},
                       headers=roles["lawyer"]).status_code == 403


def test_admin_does_admin_work_only(client):
    """Admin: dashboard, add prisoners, lawyer accounts, assignment, audit, legal data — never case contents."""
    admin, lawyer = login(client, "admin@nyayasetu.test"), login(client, "lawyer@nyayasetu.test")
    target = pid("T-DEMO-ELIG")
    p = client.get(f"/api/persons/{target}", headers=lawyer).json()
    denied = [client.get("/api/persons", headers=admin), client.get(f"/api/persons/{target}", headers=admin),
              client.get(f"/api/persons/{target}/insights", headers=admin), client.get(f"/api/persons/{target}/drafts", headers=admin),
              client.get(f"/api/documents/{p['documents'][0]['id']}", headers=admin), client.get("/api/review", headers=admin),
              client.get("/api/alerts", headers=admin), client.get("/api/dashboard/dlsa", headers=admin),
              client.post(f"/api/persons/{target}/documents", files={"file": ("x.txt", b"x")}, headers=admin)]
    for r in denied:
        assert r.status_code in (403, 404), (r.request.url, r.status_code)
        assert "Suresh" not in r.text
    reg = client.get("/api/register", headers=admin).json()
    row = next(r for r in reg if r["id"] == target)
    assert row["name"] == "Suresh Kumar" and "cases" not in row and "custody_days" not in row
    r = new_prisoner(client, admin, name="Admin Added", jail="Central Prison Kalaburagi", district="Kalaburagi")
    assert r.status_code == 200 and r.json()["intake"] == "pending_review"


def test_only_the_assigned_lawyer_uploads_documents(client):
    target = pid("T-DEMO-NOT")
    for email in ("jail@nyayasetu.test", "dlsa@nyayasetu.test", "reviewer@nyayasetu.test", "admin@nyayasetu.test"):
        r = client.post(f"/api/persons/{target}/documents", files={"file": ("x.txt", b"FIR")}, headers=login(client, email))
        assert r.status_code == 403, (email, r.status_code)
    r = client.post(f"/api/persons/{target}/documents", files={"file": ("x.txt", b"FIR")}, headers=login(client, "lawyer2@nyayasetu.test"))
    assert r.status_code == 404  # a lawyer who is not assigned
    r = client.post(f"/api/persons/{target}/documents", files={"file": ("note.txt", b"Adjourned.\n")},
                    headers=login(client, "lawyer@nyayasetu.test"))
    assert r.status_code == 200


def test_returned_intake_blocks_assignment(client):
    jail, admin, reviewer = login(client, "jail@nyayasetu.test"), login(client, "admin@nyayasetu.test"), login(client, "reviewer@nyayasetu.test")
    new_id = new_prisoner(client, jail, name="Returned Intake").json()["id"]
    item = next(i for i in client.get("/api/review", params={"kind": "new_prisoner"}, headers=reviewer).json() if i["ref_id"] == new_id)
    assert item["payload"]["arrest_date"] and item["payload"]["charges"] == ["BNS 303(2)"]
    client.post(f"/api/review/{item['id']}/resolve", json={"action": "reject", "note": "arrest date looks wrong"}, headers=reviewer)
    l1 = client.get("/api/auth/me", headers=login(client, "lawyer@nyayasetu.test")).json()["id"]
    assert client.post(f"/api/persons/{new_id}/assign", json={"lawyer_id": l1}, headers=admin).status_code == 409
    assert next(r for r in client.get("/api/register", headers=admin).json() if r["id"] == new_id)["intake"] == "returned"


def test_stale_result_is_recomputed_before_display(client, monkeypatch):
    lawyer = login(client, "lawyer@nyayasetu.test")
    target = pid("T-DEMO-NOT")
    d0 = client.get(f"/api/persons/{target}", headers=lawyer).json()
    monkeypatch.setattr(settings, "today_override", settings.today() + timedelta(days=1))
    d1 = client.get(f"/api/persons/{target}", headers=lawyer).json()
    assert d1["eligibility"]["as_of"] == settings.today().isoformat()
    assert d1["eligibility"]["cases"][0]["custody_days"] == d0["eligibility"]["cases"][0]["custody_days"] + 1


def test_typed_in_prisoner_stays_in_review_until_documents_support_it(client):
    """A new prisoner's details are typed in by hand: no definite result until documents back them up."""
    admin, lawyer = login(client, "admin@nyayasetu.test"), login(client, "lawyer@nyayasetu.test")
    today = settings.today()
    arrest, remand = today - timedelta(days=200), today - timedelta(days=199)
    r = new_prisoner(client, admin, name="Evidence Needed", jail=JAIL_A, district="Bengaluru Urban",
                     arrest_date=arrest.isoformat(), first_remand_date=remand.isoformat(),
                     offence_date=(arrest - timedelta(days=2)).isoformat(),
                     charge_sheet_date=(remand + timedelta(days=40)).isoformat())
    new_id = r.json()["id"]
    assert r.json()["status"] == "REVIEW"
    verify_intake(client, new_id)
    l1 = client.get("/api/auth/me", headers=lawyer).json()["id"]
    client.post(f"/api/persons/{new_id}/assign", json={"lawyer_id": l1}, headers=admin)
    c = client.get(f"/api/persons/{new_id}", headers=lawyer).json()["eligibility"]["cases"][0]
    assert c["status"] == "REVIEW" and c["counted_days"] == 201  # numbers shown, but only as provisional
    assert sum(f["code"] == "UNVERIFIED_MANUAL_ENTRY" for f in c["flags"]) == 3
    up = lambda name, text: client.post(f"/api/persons/{new_id}/documents", files={"file": (name, text.encode())},  # noqa: E731
                                        headers=lawyer)
    fmt = lambda d: d.strftime("%d/%m/%Y")  # noqa: E731
    up("fir.txt", f"FIRST INFORMATION REPORT\nPolice Station: Test PS   FIR No.: 12/2026\nAccused: Evidence Needed s/o Ramappa\n"
                  f"Date and time of occurrence: {fmt(arrest - timedelta(days=2))} at 21:00 hrs\nSections: u/s 303(2) BNS\n")
    # a document for SOMEONE ELSE with the right dates must not count as support
    up("other.txt", f"ARREST MEMO\nFIR No.: 12/2026   Police Station: Test PS\nName of arrested person: Mohammed Rafiq s/o Yusuf Khan\n"
                    f"Date and time of arrest: {fmt(arrest)} 10:00 hrs\n")
    assert client.get(f"/api/persons/{new_id}", headers=lawyer).json()["status"] == "REVIEW"
    up("arrest.txt", f"ARREST MEMO\nFIR No.: 12/2026   Police Station: Test PS\nName of arrested person: Evidence Needed s/o Ramappa\n"
                     f"Date and time of arrest: {fmt(arrest)} 10:00 hrs\nGrounds of arrest communicated: Yes\n")
    up("remand.txt", f"ORDER ON REMAND\nIn the Court of JMFC Test Court\nCrime No. 12/2026 of Test PS\nAccused: Evidence Needed s/o Ramappa\n"
                     f"Remanded to judicial custody from {fmt(remand)}\n")
    d = client.get(f"/api/persons/{new_id}", headers=lawyer).json()
    assert d["status"] == "NOT_YET", d["eligibility"]["cases"][0]["needs_verification"]


def test_the_same_prisoner_cannot_be_added_twice(client):
    admin, jail_other = login(client, "admin@nyayasetu.test"), None
    reviewer = login(client, "reviewer@nyayasetu.test")
    base = dict(jail=JAIL_A, district="Bengaluru Urban")
    first = new_prisoner(client, admin, name="Duplicate Check", cnr="KA01999900012026", police_station="Hebbal PS",
                         fir_number="301/2026", **base)
    assert first.status_code == 200
    # same CNR (even with different spacing/case)
    r = new_prisoner(client, admin, name="Someone Else", cnr="ka01 9999 0001 2026", **base)
    assert r.status_code == 409 and "Duplicate Check" in r.json()["detail"] and "CNR" in r.json()["detail"]
    # same FIR at the same police station
    r = new_prisoner(client, admin, name="Another Name", fir_number="301/2026", police_station="HEBBAL PS", **base)
    assert r.status_code == 409 and "FIR" in r.json()["detail"]
    # same person by name + father + DOB, spelled differently
    body_dob = "1990-05-05"
    a = client.post("/api/persons", headers=admin, json={"name": "Nagaraju Gowda", "relative_name": "Ramegowda", "relation": "s/o",
                    "dob": body_dob, "jail": JAIL_A, "district": "Bengaluru Urban",
                    "case": {"court": "JMFC", "charges": [{"act": "BNS", "section": "303(2)"}], "arrest_date": "2026-05-01"}})
    assert a.status_code == 200
    b = client.post("/api/persons", headers=admin, json={"name": "NAGARAJU GOUDA", "relative_name": "Rame Gowda", "relation": "s/o",
                    "dob": body_dob, "jail": JAIL_A, "district": "Bengaluru Urban",
                    "case": {"court": "JMFC", "charges": [{"act": "BNS", "section": "303(2)"}], "arrest_date": "2026-05-01"}})
    assert b.status_code == 409 and "date of birth" in b.json()["detail"]
    # a look-alike (same name, different father) is a different person: allowed, but sent to the reviewer
    c = client.post("/api/persons", headers=admin, json={"name": "Nagaraju Gowda", "relative_name": "Siddappa", "relation": "s/o",
                    "dob": "1975-01-01", "jail": JAIL_A, "district": "Bengaluru Urban",
                    "case": {"court": "JMFC", "charges": [{"act": "BNS", "section": "303(2)"}], "arrest_date": "2026-05-01"}})
    assert c.status_code == 200
    # jail staff of ANOTHER jail are told only that it exists — not whose record it is
    client.post("/api/admin/users", json={"email": "jail.kalaburagi@nyayasetu.test", "name": "Kalaburagi jail",
                                          "role": "jail_staff", "password": "kalaburagi-123", "jail": "Central Prison Kalaburagi"},
                headers=admin)
    other = login(client, "jail.kalaburagi@nyayasetu.test", "kalaburagi-123")
    r = new_prisoner(client, other, name="Whoever", cnr="KA01999900012026", district="Kalaburagi")
    assert r.status_code == 409 and "Duplicate Check" not in r.json()["detail"] and JAIL_A not in r.json()["detail"]
    assert "duplicate_blocked" in [x["action"] for x in audit_actions(client, admin)]
    del jail_other, reviewer
